"""產生 PPTX：16:9，每頁 GIF 鋪滿當背景，文字是 PowerPoint 文字框（微軟正黑體），
換頁 Morph（舊版退成 Fade）、換到該頁時文字自動依序淡入（一按就是下一頁）。

python src/build_pptx.py [--out 成品/<檔名>.pptx]
重用 kit 的 gen_pptx：textbox、_font、para（hangingPunct=0）、rect（圓角＋透明度）、n_lines（MSJH 量行數）。
"""
import argparse
import json
import os
import sys

from PIL import ImageFont

SRC = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SRC)
KIT = os.environ.get("SLIDEKIT_KIT", "/Users/kbchang/localcode/slidekit/pathology-slide-kit-draft")
sys.path.insert(0, SRC)
sys.path.insert(0, os.path.join(KIT, "scripts"))
import layout as L  # noqa: E402
import pptx_anim as A  # noqa: E402
import gen_pptx as gp  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

OUT_DEFAULT = os.path.join(ROOT, "成品", "讓生命乘著歌聲綻放成一道彩虹_動畫投影片_草稿.pptx")
IN = lambda px: px / L.PX_PER_IN          # px → 英寸
PT = lambda px: px / 2.0                  # px → pt
H = lambda c: c.lstrip("#")

_f = {}


def mfont(px, bold=False):
    k = (px, bold)
    if k not in _f:
        _f[k] = ImageFont.truetype(L.font_path(bold), px)
    return _f[k]


def textbox(slide, x, y, w, h, name, anchor=MSO_ANCHOR.TOP):
    tb, tf = gp.textbox(slide, IN(x), IN(y), IN(w), IN(h), anchor)
    tb.name = name
    tf.word_wrap = True
    return tb, tf


def rrect(slide, x, y, w, h, fill, alpha=None, r=36, border=None, name=None):
    adj = min(0.5, r / max(1, min(w, h)))
    shp = gp.rect(slide, IN(x), IN(y), IN(w), IN(h), H(fill), H(border) if border else None, 2.0, True, adj,
                  None if alpha is None else alpha * 100)
    if name:
        shp.name = name
    return shp


def put_text(tf, text, px, color, bold=False, align=PP_ALIGN.LEFT, line=1.0, first=True):
    return gp.para(tf, text, PT(px), H(color), bold, align, 0, line, first=first, markup=False)


def fit_px(text, want_px, max_w, bold=False):
    px = want_px
    while px > 20 and mfont(px, bold).getlength(text) > max_w:
        px -= 2
    return px


def story_slide(s, p, lines):
    C, T = L.CARD, L.TAG
    x, y, w, h = L.card_box(len(lines))
    tag = f"{p['chapter']}／8〈{p['song']}〉"
    tw = mfont(T["font_px"]).getlength(tag) + 2 * T["pad_x"] + 8
    pill = rrect(s, T["x"], T["y"], tw, T["h"], T["fill"], T["alpha"], T["h"] / 2, name="小標")
    tf = pill.text_frame
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    put_text(tf, tag, T["font_px"], T["ink"], align=PP_ALIGN.CENTER)
    card = rrect(s, x, y, w, h, C["fill"], C["alpha"], C["radius"], C["border"], name="文字卡")
    # 文字框比內寬多 0.2 英寸，換行點完全照 layout.wrap 算好的行（以 ⏎ 換行），PowerPoint 不會自己再換
    tb, tf = textbox(s, x + C["pad_x"], y + C["pad_top"] - 4, w - 2 * C["pad_x"] + 29, len(lines) * C["line_px"] + 20, "故事文字")
    line = C["line_px"] / (PT(C["font_px"]) * 2 * 1.12)   # 95px 行距 → 行距倍數（kit 校正係數 1.12）
    put_text(tf, "⏎".join(lines), C["font_px"], C["ink"], line=round(line, 3))
    return [(pill, "fade", 0, 500), (card, "fade", 300, 700), (tb, "rise", 450, 700)]


def song_slide(s, p):
    S = L.SONG
    box = rrect(s, S["x"], S["y"], S["w"], S["h"], S["fill"], S["alpha"], S["radius"], name="歌名框")
    cx, inner = S["x"], S["w"]
    lab, tf = textbox(s, cx, S["y"] + 58, inner, S["label_px"] * 1.5, "第幾首")
    put_text(tf, p["label"], S["label_px"], S["ink"], align=PP_ALIGN.CENTER)
    title = f"〈{p['song']}〉"
    tpx = fit_px(title, S["title_px"], inner - 70, True)
    tt, tf = textbox(s, cx, S["y"] + 58 + S["label_px"] * 1.5 + 30, inner, tpx * 1.45, "歌名", MSO_ANCHOR.MIDDLE)
    tf.word_wrap = False
    put_text(tf, title, tpx, S["ink"], bold=True, align=PP_ALIGN.CENTER)
    fw = mfont(S["form_px"], True).getlength(p["form"]) + 84
    fh = S["form_px"] + 36
    pill = rrect(s, cx + (inner - fw) / 2, S["y"] + S["h"] - fh - 64, fw, fh, S["form_fill"], None, fh / 2, name="演出形式")
    tf = pill.text_frame
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.word_wrap = False
    put_text(tf, p["form"], S["form_px"], S["form_ink"], bold=True, align=PP_ALIGN.CENTER)
    return [(box, "fade", 0, 500), (lab, "fade", 200, 500), (tt, "zoom", 400, 800), (pill, "fade", 1000, 500)], tpx


def cover_slide(s, p):
    V = L.COVER
    x, y, w, h = V["title_box"]
    tbox = rrect(s, x, y, w, h, V["fill"], V["alpha"], V["radius"], name="標題框")
    tpx = fit_px(p["title"], V["title_px"], w - 80, True)
    tt, tf = textbox(s, x, y, w, h, "標題", MSO_ANCHOR.MIDDLE)
    tf.word_wrap = False
    put_text(tf, p["title"], tpx, V["title_ink"], bold=True, align=PP_ALIGN.CENTER)
    x, y, w, h = V["credits_box"]
    cbox = rrect(s, x, y, w, h, V["fill"], V["alpha"], V["radius"], name="計畫名稱框")
    ct, tf = textbox(s, x, y, w, h, "計畫名稱", MSO_ANCHOR.MIDDLE)
    tf.word_wrap = False
    for i, (ln, px) in enumerate(zip(p["credits"], V["credits_px"])):
        para = put_text(tf, ln, px, V["credits_ink"], bold=px >= 40, align=PP_ALIGN.CENTER, line=1.0, first=(i == 0))
        para.space_after = Pt(8)
    return [(tbox, "fade", 0, 600), (tt, "wipe", 300, 1500), (cbox, "fade", 1200, 600), (ct, "rise", 1400, 800)]


def notes_for(i, pages):
    p = pages[i]
    nxt = pages[i + 1] if i + 1 < len(pages) else None
    if nxt is None:
        cue = "（最後一頁）"
    elif nxt["kind"] == "story" and p.get("chapter") == nxt["chapter"]:
        cue = "下一按：本章下一段故事"
    elif nxt["kind"] == "story":
        cue = f"下一按：第 {nxt['chapter']} 章〈{nxt['song']}〉的故事（新的一章）"
    elif nxt["kind"] == "song":
        cue = f"下一按：歌曲頁〈{nxt['song']}〉（{nxt['form']}）"
    else:
        cue = "下一按：結尾頁"
    if p["kind"] == "story":
        return f"第 {p['chapter']} 章〈{p['song']}〉第 {p['para']}／{p['paras']} 段\n\n{p['text']}\n\n{cue}"
    if p["kind"] == "song":
        return (f"{p['label']}〈{p['song']}〉（{p['form']}）\n背景動畫整首歌一直循環；唱完再按下一頁。\n\n{cue}\n\n本章故事：\n{p['notes']}")
    if p["kind"] == "cover":
        return f"封面。換到這頁時標題會自動擦入。\n\n{cue}"
    return "結尾頁。"


def build(out):
    story = json.load(open(os.path.join(SRC, "story.json"), encoding="utf-8"))
    pages = story["pages"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
    blank = prs.slide_layouts[6]
    inner = L.CARD["w"] - 2 * L.CARD["pad_x"]
    mf = mfont(L.CARD["font_px"])
    missing, report = [], []
    for i, p in enumerate(pages):
        s = prs.slides.add_slide(blank)
        gif = os.path.join(ROOT, "assets", p["scene"] + ".gif")
        if os.path.exists(gif):
            pic = s.shapes.add_picture(gif, 0, 0, prs.slide_width, prs.slide_height)
            pic.name = "背景"
            pic._element.nvPicPr.cNvPr.set("descr", p["alt"])
        else:
            missing.append(p["scene"])
            gp.set_bg(s, "F6F1E6")
        if p["kind"] == "story":
            lines = L.wrap(p["text"], inner, lambda t: mf.getlength(t))
            # 用 kit 的 n_lines 再量一次（PowerPoint 的寬度），每行都要是 1 行
            for ln in lines:
                n = gp.n_lines(ln, PT(L.CARD["font_px"]), IN(inner) + 0.2)
                if n != 1:
                    report.append(f"{p['id']} 有一行 PowerPoint 會換成 {n} 行：{ln}")
            eff = story_slide(s, p, lines)
            first_of_ch = p["para"] == 1
            A.transition(s, "fade" if first_of_ch else "morph", 1200 if first_of_ch else 900)
        elif p["kind"] == "song":
            eff, tpx = song_slide(s, p)
            A.transition(s, "morph", 1200)
        else:
            eff = cover_slide(s, p)
            A.transition(s, "fade", 1500 if p["kind"] == "end" else 1000)
        A.entrances(s, eff)
        s.notes_slide.notes_text_frame.text = notes_for(i, pages)
    prs.core_properties.title = story["title"]
    prs.core_properties.subject = "樂齡合唱音樂會動畫投影片（草稿）"
    os.makedirs(os.path.dirname(out), exist_ok=True)
    prs.save(out)
    mb = os.path.getsize(out) / 1e6
    print(f"{os.path.relpath(out, ROOT)}：{len(pages)} 張，{mb:.1f} MB" + (f"；還沒有 GIF：{' '.join(missing)}" if missing else ""))
    for r in report:
        print("  ⚠ " + r)
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=OUT_DEFAULT)
    a = ap.parse_args()
    build(a.out)
