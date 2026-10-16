"""文字比對：story.json、網站 index.html、PPTX、claude.ai 簡報檔的文字，和劇本原文一字不差。

python src/check_text.py [--pptx 成品/x.pptx] [--slides _archive/slides/project]
允許的差異只有（交件時列出）：「四重」顯示成「四重唱」；計畫名稱的「-」→「－」、「114 -116」→「114–116」、分行。
結束碼：有不一致就是 1。
"""
import argparse
import glob
import html as H
import json
import os
import re
import sys
import unicodedata

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "src")
NFC = lambda s: unicodedata.normalize("NFC", s)


def original():
    raw = NFC(open(os.path.join(SRC, "script_original.txt"), encoding="utf-8").read())
    lines = [l.strip() for l in raw.splitlines() if l.strip()]
    title = lines[0]
    head = re.compile(r"^(\d)\.〈(.+?)〉——(.+)$")
    chapters = []
    for l in lines[1:]:
        m = head.match(l)
        if m:
            chapters.append({"n": int(m.group(1)), "song": m.group(2), "form": m.group(3), "paras": []})
        else:
            chapters[-1]["paras"].append(l)
    return title, chapters


def credits_ok(lines):
    raw = NFC(open(os.path.join(SRC, "credits_original.txt"), encoding="utf-8").read().strip())
    norm = lambda s: re.sub(r"[\s　]", "", s).replace("－", "-").replace("–", "-")
    return norm("".join(lines)) == norm(raw)


def expected():
    title, chapters = original()
    exp = {"title": title, "story": [], "songs": []}
    for c in chapters:
        for p in c["paras"]:
            exp["story"].append(p)
        exp["songs"].append((c["song"], {"四重": "四重唱"}.get(c["form"], c["form"])))
    return exp


def compare(name, got, exp, problems):
    if got["title"] != exp["title"]:
        problems.append(f"{name} 標題不同：{got['title']!r}")
    if len(got["story"]) != len(exp["story"]):
        problems.append(f"{name} 故事段數 {len(got['story'])}，原文 {len(exp['story'])}")
    for i, (g, e) in enumerate(zip(got["story"], exp["story"])):
        if g != e:
            k = next((j for j in range(min(len(g), len(e))) if g[j] != e[j]), min(len(g), len(e)))
            problems.append(f"{name} 第 {i + 1} 段不同（第 {k + 1} 字起）：原文「{e[max(0, k - 5):k + 10]}」／這裡「{g[max(0, k - 5):k + 10]}」")
    if got["songs"] != exp["songs"]:
        problems.append(f"{name} 歌名或演出形式不同：{got['songs']}")
    for c in got.get("credits", []):
        if not credits_ok(c):
            problems.append(f"{name} 計畫名稱不同：{c}")


def from_story():
    s = json.load(open(os.path.join(SRC, "story.json"), encoding="utf-8"))
    P = s["pages"]
    return {"title": s["title"], "story": [p["text"] for p in P if p["kind"] == "story"],
            "songs": [(p["song"], p["form"]) for p in P if p["kind"] == "song"],
            "credits": [p["credits"] for p in P if p["kind"] in ("cover", "end")]}


def from_site():
    t = open(os.path.join(ROOT, "index.html"), encoding="utf-8").read()
    data = json.loads(re.search(r'<script id="data" type="application/json">(.*?)</script>', t, re.S).group(1).replace("<\\/", "</"))
    P = data["pages"]
    return {"title": next(p["title"] for p in P if p["k"] == "cover"),
            "story": ["".join(p["lines"]) for p in P if p["k"] == "story"],
            "songs": [(p["song"], p["form"]) for p in P if p["k"] == "song"],
            "credits": [p["credits"] for p in P if p["k"] in ("cover", "end")],
            "story_full": [p["text"] for p in P if p["k"] == "story"]}


def from_pptx(path):
    from pptx import Presentation
    prs = Presentation(path)
    got = {"title": None, "story": [], "songs": [], "credits": []}
    for s in prs.slides:
        names = {sh.name: sh for sh in s.shapes}
        kind = "story" if "故事文字" in names else ("song" if "歌名" in names else ("cover" if "標題" in names else None))
        def txt(n):
            sh = names.get(n)
            return NFC(sh.text_frame.text.replace("\x0b", "").replace("\n", "")) if sh is not None and sh.has_text_frame else None
        if kind == "story":
            got["story"].append(txt("故事文字"))
        elif kind == "song":
            got["songs"].append((txt("歌名").strip("〈〉"), txt("演出形式")))
        elif kind in ("cover", "end"):
            got["title"] = got["title"] or txt("標題")
            sh = names.get("計畫名稱")
            got["credits"].append([NFC(p.text) for p in sh.text_frame.paragraphs])
    return got


def from_slides(root):
    files = sorted(glob.glob(os.path.join(root, "slides", "*.html")))
    deck = json.load(open(os.path.join(root, "deck.json"), encoding="utf-8"))
    order = deck["order"]
    by = {os.path.splitext(os.path.basename(f))[0]: open(f, encoding="utf-8").read() for f in files}
    got = {"title": None, "story": [], "songs": [], "credits": []}
    strip = lambda s: NFC(H.unescape(re.sub(r"<[^>]+>", "", s)))
    for sid in order:
        t = by[sid]
        kind = re.search(r'data-kind="(\w+)"', t).group(1)
        if kind == "story":
            got["story"].append(strip(re.search(r'<p[^>]*data-role="text"[^>]*>(.*?)</p>', t, re.S).group(1)).replace("\n", ""))
        elif kind == "song":
            got["songs"].append((strip(re.search(r'<h2[^>]*data-role="song"[^>]*>(.*?)</h2>', t, re.S).group(1)).strip("〈〉"),
                                 strip(re.search(r'<p[^>]*data-role="form"[^>]*>(.*?)</p>', t, re.S).group(1))))
        else:
            got["title"] = got["title"] or strip(re.search(r'<h1[^>]*>(.*?)</h1>', t, re.S).group(1))
            got["credits"].append([strip(x) for x in re.findall(r'<p[^>]*data-role="credit"[^>]*>(.*?)</p>', t, re.S)])
    return got


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pptx")
    ap.add_argument("--slides")
    a = ap.parse_args()
    exp = expected()
    problems = []
    compare("story.json", from_story(), exp, problems)
    if os.path.exists(os.path.join(ROOT, "index.html")):
        site = from_site()
        compare("網站", site, exp, problems)
        if site["story_full"] != exp["story"]:
            problems.append("網站手機版全文和原文不同")
    if a.pptx:
        compare("PPTX", from_pptx(a.pptx), exp, problems)
    if a.slides:
        compare("Slides", from_slides(a.slides), exp, problems)
    checked = ["story.json", "網站"] + (["PPTX"] if a.pptx else []) + (["Slides"] if a.slides else [])
    print(f"比對：{'、'.join(checked)}；原文 {len(exp['story'])} 段、{len(exp['songs'])} 首歌")
    print("結果：" + ("一字不差（只有列出的小調整）" if not problems else "\n  " + "\n  ".join(problems)))
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
